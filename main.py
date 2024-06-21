import openpyxl
from openpyxl.chart import Reference, Series, PieChart, BarChart,LineChart, ScatterChart
import numpy as np
import random

# Creating a virtual workbook to store the data
wb = openpyxl.Workbook() 

# Making the sheet active for entering data
sheet = wb.active 

#for value in range(1,11):
  #sheet["A"+str(value)] = value


#refObj = openpyxl.chart.Reference(sheet, min_col=1, min_row=1, max_col=1,max_row=10)
#seriesObj=openpyxl.chart.Series(refObj, title="Monthly Sales Report")

# Data to be added into virtual workbook
course = ["Business Analytics", "Health Analytics", "IT Business Analytics", "Business Administration", "Crime Analytics", "Marketing"]

population = [45, 36, 50, 25, 18, 40]

#Generating squares of numbers 
x_values = [i for i in range(11)]
y_values = [i**2 for i in range(11)]

# Random generation of numbers for scatter plot
Xvalues = [random.uniform(1, 100) for i in range (100)]
Yvalues = [random.uniform(1, 100) for i in range (100)]

#Setting headers for individual columns
sheet["A1"] = "Course "
sheet["B1"] = "Population"
sheet["D1"] = "X_values"
sheet["E1"] = "Y_values"
sheet["G1"] = "Xvalues"
sheet["H1"] = "Yvalues"


#Adding data to individual columns
for number in range(len(course)):
  sheet["A" + str(number+2)] = course[number]
  sheet["B" + str(number+2)] = population[number]


for number in range(len(x_values)):
  sheet["D" + str(number+2)] = x_values[number]
  sheet["E" + str(number+2)] = y_values[number]


for number in range(len(Yvalues)):
  sheet["G" + str(number+2)] = Xvalues[number]
  sheet["H" + str(number+2)] = Yvalues[number]


# Individual data for the respective plot
data = Reference(sheet, min_col=2, min_row=1, max_col=2, max_row=len(course) + 1)

data2 = Reference(sheet, min_col=4, min_row=1, max_col=5, max_row=len(x_values) + 1)

data3 = Reference(sheet, min_col=8, min_row=1, max_col=8, max_row=len(Yvalues) + 1)

label = Reference(sheet, min_col=1, min_row=2, max_row=len(course) + 1)

label2 = Reference(sheet, min_col=4, min_row=2, max_row=len(x_values) + 1)

label3 = Reference(sheet, min_col=7, min_row=2, max_row=len(Yvalues) + 1)


#Pie chart
pc = PieChart()
pc.add_data(data, titles_from_data=True)
pc.set_categories(label)
pc.title = "Piechart of various programs"
sheet.add_chart(pc, "J1")

#Barchart
bc = BarChart()
bc.add_data(data, titles_from_data=True)
bc.set_categories(label)
bc.title = "Barchart distribution of various programs"
sheet.add_chart(bc, "J16")

#Linechart
lc = LineChart()
lc.add_data(data2, titles_from_data=True)
lc.set_categories(label2)
lc.title = "Linechart of squares of Numbers"
sheet.add_chart(lc, "J32")


#Scatter
sc = ScatterChart()
sc.series.append(Series(data3, xvalues=label3, title="Scatter"))
#sc.set_categories(label3)
sc.title = "Scatterchart"
sc.x_axis.title = "X values"
sc.y_axis.title = "Y values"
sc.series[0].marker.symbol = "circle"
sc.series[0].graphicalProperties.line.noFill = True
sheet.add_chart(sc, "J50")


# Statistical analysis of the data
mean = np.mean(population)
median = np.median(population)
std = np.std(population)

sheet["A10"] = f"The mean is {round(mean)}"
sheet["A11"] = f"The median is {round(median)}"
sheet["A12"] = f"The standard deviation is {round(std)}"


#chartObj = openpyxl.chart.BarChart()
#chartObj = openpyxl.chart.LineChart()
#chartObj = openpyxl.chart.ScatterChart()
#chartObj = openpyxl.chart.PieChart()


#chartObj.title = "My Chart"
#chartObj.append(seriesObj)

#sheet.add_chart(chartObj, "C5")

#Saving chart and data to the workbook 
wb.save('Excel/OpenpyxlCharts.xlsx')